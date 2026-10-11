"""把每页图片打包成全屏图片版 PPTX。"""

import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def export_pptx(project: dict, pdir: Path) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    for s in project["slides"]:
        if not s.get("image"):
            continue
        slide = prs.slides.add_slide(blank)
        img_bytes = _cover_crop(pdir / s["image"])
        slide.shapes.add_picture(img_bytes, Emu(0), Emu(0), SLIDE_W, SLIDE_H)
        notes = s["title"] + ("\n" + s["subtitle"] if s["subtitle"] else "")
        slide.notes_slide.notes_text_frame.text = notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _cover_crop(path: Path) -> io.BytesIO:
    """居中裁成精确的 16:9，避免图片在幻灯片上被拉伸。"""
    img = Image.open(path).convert("RGB")
    target = 16 / 9
    w, h = img.size
    if abs(w / h - target) > 0.005:
        if w / h > target:
            nw = round(h * target)
            img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = round(w / target)
            img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=92)
    buf.seek(0)
    return buf
